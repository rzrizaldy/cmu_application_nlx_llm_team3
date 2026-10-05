"""Part E cost/benefit inputs, computed only from measured numbers plus assumptions you enter.

  python cost.py review-sheet <run> [--n 20]
      -> labels/review_a2.xlsx: n API outputs from <run> next to their records. Time yourself
         (Timing sheet) and mark each output usable as-is (yes/no). This gives review seconds per output.
  python cost.py model <run> [<run> ...]
      -> results/cost_model.json, using:
         runs/<run>/metrics.json           exact match, tokens, latency per request (measured)
         results/partd_metrics.json        guardrail added latency/tokens (measured), if present
         labels/review_timing.json         review seconds per API output (measured by review-sheet)
         data/cost_assumptions.json        $ per staff hour, $ per compute hour, setup hours, volume, and the
                                           manual seconds per record without the API (yours, each with a source)

For each run, per record:
  manual_seconds        = estimated time for one record done by hand in the organization
                          (brief, Part E: "estimate how long the same task takes without the API";
                          state its basis in manual_seconds_source)
  api_human_seconds     = review_seconds + (1 - usable_rate) * manual_seconds
                          (every output is reviewed, and the unusable ones are redone by hand)
  api_compute_seconds   = mean latency (+ guardrail latency if the run is guarded)
  saving_per_record_$   = manual_$ - (api_human_$ + api_compute_$)
  break_even_volume     = setup_$ / saving_per_record_$   (records; none if saving <= 0)
"""
import json
import sys
from datetime import datetime

from a2common import DATA, LABELS, RESULTS, RUNS, corpus, content, save_json

ASSUMPTIONS = DATA / "cost_assumptions.json"
TEMPLATE = {
    "staff_cost_per_hour_usd": None, "staff_cost_source": "",
    "compute_cost_per_hour_usd": None, "compute_cost_source": "",
    "setup_hours": None, "setup_hours_source": "",
    "annual_record_volume": None, "annual_volume_source": "",
    "manual_seconds_per_record": None, "manual_seconds_source": "",
}


def review_sheet(run, n=20):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    rows = [json.loads(x) for x in (RUNS / run / "responses.jsonl").read_text().splitlines() if x.strip()][:n]
    text = {r["doc_id"]: content(r) for r in corpus()}
    path = LABELS / "review_a2.xlsx"
    if path.exists():
        raise SystemExit(f"{path} exists; not overwriting")
    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    ws.append(["row", "doc_id", "record", "api_output", "usable_as_is (yes/no)", "notes"])
    for i, r in enumerate(rows, 1):
        out = r.get("parsed") if r.get("parsed") is not None else r.get("text")
        ws.append([i, r["id"], text[r["id"]], json.dumps(out, ensure_ascii=False) if not isinstance(out, str) else out, "", ""])
    for col, w in zip("ABCDEF", [5, 26, 80, 50, 18, 30]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=4):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    t = wb.create_sheet("Timing")
    t.append(["run", "rows", "start_time (HH:MM)", "end_time (HH:MM)", "date (YYYY-MM-DD)"])
    t.append([run, len(rows), "", "", ""])
    wb.save(path)
    print(f"wrote {path} ({len(rows)} outputs from {run})")


def review_timing():
    from openpyxl import load_workbook
    path = LABELS / "review_a2.xlsx"
    wb = load_workbook(path, data_only=True)
    run, n, start, end, date = list(wb["Timing"].values)[1]
    fmt = "%Y-%m-%d %H:%M"
    if not all((start, end, date)):
        raise SystemExit("Review Timing sheet needs real start time, end time and date")
    secs = (datetime.strptime(f"{str(date)[:10]} {str(end)[:5]}", fmt) - datetime.strptime(f"{str(date)[:10]} {str(start)[:5]}", fmt)).total_seconds()
    if secs <= 0:
        raise SystemExit("Review Timing end must be after start")
    marks = [str(r[4]).strip().lower() for r in list(wb["Review"].values)[1:] if r[1]]
    if len(marks) != n or any(m not in {"yes", "no"} for m in marks):
        raise SystemExit("Mark every review output exactly yes or no before cost modeling")
    out = {"run": run, "rows": n, "seconds_per_output": round(secs / n, 1),
           "human_usable_rate": sum(m == "yes" for m in marks) / len(marks) if marks else None}
    save_json(LABELS / "review_timing.json", out)
    return out


def model(runs):
    if not ASSUMPTIONS.exists():
        save_json(ASSUMPTIONS, TEMPLATE)
        raise SystemExit(f"Fill in {ASSUMPTIONS} (values and sources) and rerun.")
    a = json.loads(ASSUMPTIONS.read_text())
    missing = [k for k, v in a.items() if v is None]
    if missing:
        raise SystemExit(f"{ASSUMPTIONS}: fill {missing}")
    manual_s = a.get("manual_seconds_per_record")
    if not isinstance(manual_s, (int, float)) or manual_s <= 0 or not str(a.get("manual_seconds_source", "")).strip():
        raise SystemExit(f"{ASSUMPTIONS}: set manual_seconds_per_record (> 0) and manual_seconds_source")
    review = review_timing() if (LABELS / "review_a2.xlsx").exists() else None
    if review is None:
        raise SystemExit("Run `cost.py review-sheet <run>` and fill labels/review_a2.xlsx first.")
    partd = json.loads((RESULTS / "partd_metrics.json").read_text()) if (RESULTS / "partd_metrics.json").exists() else None
    staff, compute = a["staff_cost_per_hour_usd"] / 3600, a["compute_cost_per_hour_usd"] / 3600
    setup_usd = a["setup_hours"] * a["staff_cost_per_hour_usd"]
    table = {"assumptions": a, "manual_seconds_per_record": manual_s,
             "manual_seconds_source": a["manual_seconds_source"], "review": review, "runs": {}}
    for run in runs:
        m = json.loads((RUNS / run / "metrics.json").read_text())
        reviewed_run = review["run"] == run
        usable = review["human_usable_rate"] if reviewed_run else m["exact_match"]
        compute_s = m.get("latency_s", {}).get("mean", 0.0)
        guard_s = partd["cost_on_benign_allowed"]["added_latency_s_mean"] if partd and "guard" in run.lower() and partd.get("cost_on_benign_allowed") else 0.0
        human_s = review["seconds_per_output"] + (1 - usable) * manual_s
        manual_usd = manual_s * staff
        api_usd = human_s * staff + (compute_s + guard_s) * compute
        saving = manual_usd - api_usd
        table["runs"][run] = {
            "usable_rate": usable, "usable_rate_source": "timed_human_review" if reviewed_run else "exact_match_proxy",
            "tokens_per_request": m.get("tokens", {}).get("total_per_request"),
            "compute_seconds_per_record": round(compute_s + guard_s, 3),
            "human_seconds_per_record_with_api": round(human_s, 1),
            "manual_usd_per_record": round(manual_usd, 4), "api_usd_per_record": round(api_usd, 4),
            "saving_usd_per_record": round(saving, 4),
            "break_even_records": round(setup_usd / saving) if saving > 0 else None,
            "annual_net_saving_usd": round(saving * a["annual_record_volume"] - setup_usd, 2),
        }
    save_json(RESULTS / "cost_model.json", table)
    print(json.dumps(table["runs"], indent=1))


if __name__ == "__main__":
    cmd, args = sys.argv[1], sys.argv[2:]
    if cmd == "review-sheet":
        review_sheet(args[0], int(args[2]) if len(args) > 2 else 20)
    elif cmd == "model":
        model(args)
