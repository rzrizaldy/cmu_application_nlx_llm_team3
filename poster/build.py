"""Build the Pittsburgh 311 poster as a single A3 landscape slide (.pptx) for Google Slides."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from PIL import ImageFont
from lxml import etree
import sys

# ---------- palette & fonts ----------
INK = "111418"; TEXT = "1F2630"; MUTED = "55606C"; GOLD = "F2B705"; GOLD_DK = "7A5800"
LINE = "D3D8DE"; BG = "E8EBEE"; TINT = "F2F4F6"; TRACK = "E3E7EB"; GREY = "B6C0CC"; SLATE = "5C6B7E"
WHITE = "FFFFFF"; PALE = "D3D8DE"
SANS = "IBM Plex Sans"; COND = "IBM Plex Sans Condensed"; MONO = "IBM Plex Mono"

SUP = "/System/Library/Fonts/Supplemental/"
MPL = "/opt/anaconda3/lib/python3.13/site-packages/matplotlib/mpl-data/fonts/ttf/"
MEASURE = {  # (font file, width fudge to approximate the Plex family)
    (SANS, False): (SUP + "Arial.ttf", 1.14),
    (SANS, True): (SUP + "Arial Bold.ttf", 1.15),
    (COND, False): (SUP + "Arial Narrow.ttf", 1.14),
    (COND, True): (SUP + "Arial Narrow Bold.ttf", 1.15),
    (MONO, False): (MPL + "DejaVuSansMono.ttf", 1.0),
    (MONO, True): (MPL + "DejaVuSansMono-Bold.ttf", 1.0),
}
_fc = {}
def tw(text, font, size, bold):
    """Text width in inches."""
    path, k = MEASURE[(font, bold)]
    key = (path, size)
    if key not in _fc:
        _fc[key] = ImageFont.truetype(path, size * 10)
    return _fc[key].getlength(text) / 10 / 72 * k

LH = 1.3  # line height multiple

def lines_needed(runs, w, size):
    """runs: list of (text, font, bold). Greedy word wrap; returns number of lines (handles \n)."""
    total = 0
    # split into hard lines
    hard = [[]]
    for t, f, b in runs:
        parts = t.split("\n")
        for i, p in enumerate(parts):
            if i > 0:
                hard.append([])
            hard[-1].append((p, f, b))
    for hl in hard:
        words = []
        for t, f, b in hl:
            for i, wd in enumerate(t.split(" ")):
                words.append((wd, f, b, i > 0 or (words and t.startswith(" "))))
        n, cur = 1, 0.0
        for wd, f, b, _ in words:
            ww = tw(wd, f, size, b)
            sp = tw(" ", f, size, b) if cur > 0 else 0
            if cur > 0 and cur + sp + ww > w:
                n += 1; cur = ww
            else:
                cur += sp + ww
        total += n
    return total

# ---------- presentation ----------
W, H = 16.54, 11.69  # A3 landscape, inches
prs = Presentation()
prs.slide_width = Inches(W); prs.slide_height = Inches(H)
slide = prs.slides.add_slide(prs.slide_layouts[6])
SP = slide.shapes
WARN = []

def rgb(h): return RGBColor.from_string(h)

def box(x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE, name=None):
    s = SP.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill: s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else: s.fill.background()
    if line: s.line.color.rgb = rgb(line); s.line.width = Pt(lw)
    else: s.line.fill.background()
    s.shadow.inherit = False
    if name: s.name = name
    s.text_frame.text = ""
    return s

def text(x, y, w, paras, size=9, color=TEXT, font=SANS, bold=False, align=PP_ALIGN.LEFT,
         h=None, anchor=MSO_ANCHOR.TOP, space_after=0, spacing=LH, name=None, charsp=None):
    """paras: str | list of paragraphs; paragraph = str | list of runs; run = str | dict(text, bold, color, font, size)."""
    if isinstance(paras, str): paras = [paras]
    # measure
    nlines, extra = 0, 0
    norm = []
    for p in paras:
        runs = [p] if isinstance(p, (str, dict)) else p
        runs = [r if isinstance(r, dict) else {"text": r} for r in runs]
        norm.append(runs)
        psize = max(r.get("size", size) for r in runs)
        nl = lines_needed([(r["text"], r.get("font", font), r.get("bold", bold)) for r in runs], w, psize)
        nlines += nl
        extra += nl * psize * spacing / 72
    need = extra + space_after / 72 * (len(paras) - 1) + 0.02
    if h is None: h = need
    elif need > h + 0.01:
        WARN.append(f"overflow {name or norm[0][0]['text'][:40]!r}: need {need:.2f} have {h:.2f}")
    tb = SP.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name: tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, runs in enumerate(norm):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        psize = max(r.get("size", size) for r in runs)
        para.line_spacing = Pt(psize * spacing)
        if i < len(norm) - 1: para.space_after = Pt(space_after)
        for r in runs:
            run = para.add_run(); run.text = r["text"]
            f = run.font
            f.name = r.get("font", font); f.size = Pt(r.get("size", size))
            f.bold = r.get("bold", bold); f.color.rgb = rgb(r.get("color", color))
            if charsp or r.get("charsp"):
                run._r.get_or_add_rPr().set("spc", str(int((r.get("charsp") or charsp) * 100)))
    return h

def arrow(x1, y1, x2, y2, color=SLATE, lw=1.5):
    c = SP.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(lw)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd")); tail.set("type", "triangle"); tail.set("w", "med"); tail.set("len", "med")
    return c

def hline(x1, y, x2, color=LINE, lw=0.75):
    c = SP.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y), Inches(x2), Inches(y))
    c.line.color.rgb = rgb(color); c.line.width = Pt(lw)
    return c

def vline(x, y1, y2, color=INK, lw=1):
    c = SP.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y1), Inches(x), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(lw)
    return c

# ---------- background, header, footer ----------
box(0, 0, W, H, fill=BG, name="Background")
HEAD_H = 1.52
box(0, 0, W, HEAD_H, fill=INK, name="Header band")
MX = 0.42
text(MX, 0.18, 10.2, "CARNEGIE MELLON UNIVERSITY  ·  NL(X) AND LLM  ·  GROUP 3", size=8, font=MONO, bold=True, color=GOLD, charsp=1.2)
text(MX, 0.38, 10.4, "Can a small local model route Pittsburgh 311 requests\ncorrectly on the first try?",
     size=24, font=COND, bold=True, color=WHITE, spacing=1.05, name="Title")
text(MX, 1.2, 10.2, "Afaq  ·  Rutomo  ·  Mahika  ·  Mingchin", size=11, color=PALE, bold=False)

# hero stat
hx = 11.35
text(hx, 0.16, 1.6, "0%", size=40, font=COND, bold=True, color=GREY, spacing=1.0, h=0.62, name="Hero T0")
text(hx, 0.78, 1.6, "Prompt only (T0)", size=8.5, color=PALE)
arrow(hx + 1.35, 0.49, hx + 1.95, 0.49, color=GOLD, lw=3)
text(hx + 2.15, 0.16, 2.4, "82%", size=40, font=COND, bold=True, color=GOLD, spacing=1.0, h=0.62, name="Hero T2")
text(hx + 2.15, 0.78, 2.4, "Grounded in the city's codebook (T2)", size=8.5, color=PALE)
text(hx, 1.1, W - MX - hx, "SHARE ROUTED CORRECTLY, 50 HELD-OUT INPUTS  ·  T2 95% CI 70–92%",
     size=6.5, font=MONO, bold=True, color=GREY, charsp=0.5)

FOOT_H = 0.48
FY = H - FOOT_H
box(0, FY, W, FOOT_H, fill=INK, name="Footer band")
text(MX, FY, 12.4, [[{"text": "Takeaway.  ", "color": GOLD, "bold": True},
                     {"text": "Grounded in Pittsburgh's own codebook, a 4-billion-parameter model running on a laptop routed 82% of held-out 311 complaints correctly. Without that grounding it routed none.", "color": WHITE}]],
     size=10, h=FOOT_H, anchor=MSO_ANCHOR.MIDDLE, name="Takeaway")
text(W - MX - 3.2, FY, 3.2, "github.com/rzrizaldy/cmu_application_nlx_llm_team3", size=6.5, font=MONO, color=GREY,
     h=FOOT_H, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)

# ---------- columns ----------
GAP = 0.15
TOP = HEAD_H + 0.15
BOT = FY - 0.15
NCOL = 4
CW = (W - 2 * MX - (NCOL - 1) * GAP) / NCOL
PAD = 0.15
IW = CW - 2 * PAD

def col_x(i): return MX + i * (CW + GAP)

# Each panel = function(x, y) -> draws content, returns content height. We draw a card first with
# a provisional height and fix it afterwards (card is the first shape, so it stays behind).
def panel(ci, y, num, label, title, body_fn, fill_to=None):
    x = col_x(ci)
    card = box(x, y, CW, 1, fill=WHITE, line=LINE, lw=0.75, name=f"Panel {num}")
    cy = y + PAD
    text(x + PAD, cy, IW, f"{num}  ·  {label}", size=7, font=MONO, bold=True, color=GOLD_DK, charsp=1.0)
    cy += 0.2
    cy += text(x + PAD, cy, IW, title, size=14, font=COND, bold=True, color=INK, spacing=1.1) + 0.08
    cy = body_fn(x + PAD, cy)
    h = cy + PAD - y
    if fill_to is not None:
        if h > fill_to - y + 0.005: WARN.append(f"panel {num} overflows column by {h - (fill_to - y):.2f} in")
        h = fill_to - y
    card.height = Inches(h)
    return y + h

def para(x, y, s, size=8.6, color=TEXT, w=None, **kw):
    return y + text(x, y, w or IW, s, size=size, color=color, **kw)

B = lambda t: {"text": t, "bold": True}

# ===== Column 1 =====
def p_problem(x, y):
    y = para(x, y, "When a 311 request reaches the wrong department, it waits in that queue until someone reassigns it. The City's own guide estimates that residents misclassify 5–10% of web requests. Our team brief calls the alternative first-time-right intake:") + 0.1
    steps = ["Classify the request", "Collect the details staff need", "Ask one targeted clarification question", "Route it to the responsible department"]
    for i, s in enumerate(steps):
        c = box(x, y, 0.24, 0.24, fill=INK, shape=MSO_SHAPE.OVAL)
        text(x, y, 0.24, str(i + 1), size=8, font=MONO, bold=True, color=GOLD, h=0.24, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(x + 0.34, y, IW - 0.34, s, size=8.6, h=0.24, anchor=MSO_ANCHOR.MIDDLE)
        y += 0.3
    y += 0.04
    y = para(x, y, [[B("Our question. "), "Can a 4-billion-parameter model running on a laptop do this reliably, and what does it need to get there?"]])
    return y

def p_data(x, y):
    # stat row
    stats = [("854", "records in our four Assignment 1 corpora"), ("127", "issues with historical resolution times"), ("534 / 50", "DEV / held-out EVAL inputs")]
    sw = (IW - 2 * 0.1) / 3
    sh = 0.78
    for i, (n, l) in enumerate(stats):
        sx = x + i * (sw + 0.1)
        box(sx, y, sw, sh, fill=TINT)
        text(sx + 0.08, y + 0.07, sw - 0.16, n, size=15 if i < 2 else 13, font=COND, bold=True, color=INK, spacing=1.0, h=0.28)
        text(sx + 0.08, y + 0.37, sw - 0.16, l, size=6.8, color=MUTED, spacing=1.2, h=0.38)
    y += sh + 0.14
    # table
    cols = [("Subtopic", 1.62, PP_ALIGN.LEFT), ("Member", 0.62, PP_ALIGN.LEFT), ("Records", 0.52, PP_ALIGN.RIGHT), ("Requests", IW - 2.76, PP_ALIGN.RIGHT)]
    rows = [("Streets and Mobility", "Afaq", "190", "244,708"), ("Waste and Neighborhood", "Rutomo", "180", "225,001"),
            ("Buildings and Construction", "Mahika", "260", "74,628"), ("Parks and Public Facilities", "Mingchin", "224", "67,885"),
            ("All four subtopics", "", "854", "612,222")]
    rh = 0.22
    def row(vals, yy, bold=False, color=TEXT, font=SANS, size=7.6):
        cx = x
        for (nm, cw, al), v in zip(cols, vals):
            text(cx, yy, cw, v, size=size, bold=bold, color=color, font=font, h=rh, align=al, anchor=MSO_ANCHOR.MIDDLE)
            cx += cw
    row([c[0].upper() for c in cols], y, bold=True, color=MUTED, font=MONO, size=6.2)
    y += rh; hline(x, y, x + IW, color=INK, lw=0.75)
    for i, r in enumerate(rows):
        last = i == len(rows) - 1
        if last: hline(x, y, x + IW, color=INK, lw=0.75)
        row(r, y, bold=last); y += rh
        if not last and i < len(rows) - 2: hline(x, y, x + IW)
    y += 0.06
    y = para(x, y, "612,222 of the 815,417 WPRDC requests we joined to the codebook fall in our four subtopics.", size=7.2, color=MUTED) + 0.14
    # days to close chart
    text(x, y, IW, "Days to close, busiest issue in each subtopic", size=8.6, bold=True, color=INK); y += 0.2
    # legend
    box(x, y + 0.05, 0.12, 0.08, fill=INK); text(x + 0.17, y, 0.7, "median", size=7, color=MUTED, h=0.18, anchor=MSO_ANCHOR.MIDDLE)
    box(x + 0.75, y + 0.05, 0.12, 0.08, fill=GREY); text(x + 0.92, y, 1.0, "90th percentile", size=7, color=MUTED, h=0.18, anchor=MSO_ANCHOR.MIDDLE)
    y += 0.26
    data = [("Potholes", "Streets", 10.1, 86.8), ("Weeds/Debris", "Waste", 19.9, 128.1),
            ("Building Maintenance", "Buildings", 28.8, 383.9), ("Overgrowth", "Parks", 21.6, 161.1)]
    lw_ = 1.36; bx = x + lw_; bw = IW - lw_ - 0.42; mx = 383.9
    for nm, sub, med, p90 in data:
        text(x, y, lw_ - 0.05, nm, size=7.4, bold=True, color=INK, h=0.17, spacing=1.15)
        text(x, y + 0.16, lw_ - 0.05, sub, size=6.6, color=MUTED, h=0.15, spacing=1.15)
        for j, (v, c) in enumerate([(med, INK), (p90, GREY)]):
            by = y + 0.03 + j * 0.15
            box(bx, by, max(bw * v / mx, 0.02), 0.11, fill=c)
            text(bx + bw * v / mx + 0.04, by - 0.03, 0.42, f"{v:.1f}", size=6.8, font=MONO, color=TEXT, h=0.17, anchor=MSO_ANCHOR.MIDDLE)
        y += 0.36
    y += 0.02
    y = para(x, y, "Our tickets report these historical ranges. Days to close record a status change, which can come before the problem is fixed.", size=7.4, color=MUTED)
    return y

y = panel(0, TOP, "01", "PROBLEM", "Why first-time-right intake matters", p_problem)
panel(0, y + GAP, "02", "DATA", "We built on 815,417 historical requests", p_data, fill_to=BOT)

# ===== Column 2 =====
def p_system(x, y):
    bw = IW - 0.98
    nodes = [("Resident complaint", None), ("Guardrail", "T3 only"), ("TF-IDF retrieval", "knowledge base + DEV neighbors"),
             ("Neighbor vote on issue", None), ("Phi-4-mini", "local, greedy decoding"),
             ("Snap to codebook", "category, issue, card")]
    nh = 0.29; ag = 0.13
    for i, (t, sub) in enumerate(nodes):
        fill = TINT
        box(x, y, IW, nh, fill=fill, line=LINE)
        runs = [[{"text": t, "bold": True, "color": INK}] + ([{"text": "  " + sub, "color": MUTED, "size": 6.8}] if sub else [])]
        text(x + 0.1, y, IW - 0.2, runs, size=8.4, h=nh, anchor=MSO_ANCHOR.MIDDLE)
        y += nh
        arrow(x + bw / 2, y, x + bw / 2, y + ag, color=SLATE, lw=1.25)
        y += ag
    th = 0.72
    box(x, y, bw, th, fill=INK)
    text(x + 0.1, y + 0.07, bw - 0.2, [[{"text": "Ticket311 JSON", "bold": True, "color": GOLD}],
                                       [{"text": "domain, category, issue, department, missing details, one clarification question, confidence, abstention flag, historical resolution range", "color": PALE, "size": 6.9}]],
         size=8.4, h=th - 0.1, spacing=1.2)
    # operational evidence side box
    ox = x + bw + 0.28; ow = IW - bw - 0.28
    box(ox, y + 0.12, ow, 0.62, fill=WHITE, line=SLATE, lw=1)
    text(ox + 0.06, y + 0.12, ow - 0.12, [[{"text": "Operational evidence", "bold": True, "color": INK}], [{"text": "days to close", "color": MUTED, "size": 6.6}]],
         size=7.4, h=0.62, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, spacing=1.15)
    arrow(ox, y + 0.43, x + bw, y + 0.43, color=SLATE, lw=1.25)
    y += th + 0.12
    y = para(x, y, [["We expose the pipeline as ", {"text": "route_complaint(text)", "font": MONO, "size": 7.6}, ", built on Rutomo's LLMBox fork. It runs Phi-4-mini-instruct on an Apple M4 with greedy decoding, no repetition penalty and at most 320 new tokens."]], size=7.6, color=MUTED)
    return y

def p_designs(x, y):
    rows = [("T0", "Prompt only", "The allowed categories and the complaint"),
            ("T1", "+ Retrieval", "T0 plus three retrieved knowledge documents"),
            ("T2", "+ Tools and vote", "Three labeled DEV neighbors and the vote's top issue, snapped to its codebook card"),
            ("T3", "+ Guardrail", "T2 behind input and output filters"),
            ("T4", "LoRA finetuned", "The T0 prompt on a LoRA-finetuned Phi (269 rows, 2 epochs)")]
    lw_ = 1.18
    for i, (k, nm, d) in enumerate(rows):
        hl = k == "T2"
        dh = max(0.2, lines_needed([(d, SANS, False)], IW - lw_, 7.8) * 7.8 * 1.25 / 72)
        rh = max(dh, 0.3) + 0.08
        if hl: box(x - 0.06, y, IW + 0.12, rh, fill="FBF0C8")
        text(x, y + 0.04, 0.3, k, size=8.6, font=MONO, bold=True, color=GOLD_DK if hl else INK, h=0.18)
        text(x + 0.32, y + 0.04, lw_ - 0.36, nm, size=8.2, bold=True, color=INK, h=rh - 0.08, spacing=1.2)
        text(x + lw_, y + 0.04, IW - lw_, d, size=7.8, color=TEXT, h=rh - 0.07, spacing=1.25)
        y += rh
        if i < len(rows) - 1: hline(x, y, x + IW)
    y += 0.14
    stats = [("50", "held-out EVAL inputs, the same for every design"), ("0", "EVAL items found in the knowledge base or DEV"), ("1,000", "bootstrap samples per 95% interval")]
    sw = (IW - 0.2) / 3
    for i, (n, l) in enumerate(stats):
        sx = x + i * (sw + 0.1)
        box(sx, y, sw, 0.74, fill=TINT)
        text(sx + 0.08, y + 0.07, sw - 0.16, n, size=15, font=COND, bold=True, color=INK, h=0.28, spacing=1.0)
        text(sx + 0.08, y + 0.35, sw - 0.16, l, size=6.8, color=MUTED, h=0.36, spacing=1.2)
    y += 0.74
    return y

y = panel(1, TOP, "03", "SYSTEM", "Our intake and routing pipeline", p_system)
panel(1, y + GAP, "04", "METHOD", "Five system designs on one test set", p_designs, fill_to=BOT)

# ===== Column 3 =====
def p_result(x, y):
    y = para(x, y, "Routed correctly means category and department are both right. Whiskers are 95% bootstrap intervals; those for T0, T1 and T2 do not overlap.", size=7.6, color=MUTED) + 0.12
    ax_w = 0.36; px = x + ax_w; pw = IW - ax_w; ph = 1.32
    top = y + 0.08
    for k, v in enumerate([0, 25, 50, 75, 100]):
        gy = top + ph * (1 - v / 100)
        hline(px, gy, px + pw, color=INK if v == 0 else TRACK, lw=0.75 if v else 1)
        text(x, gy - 0.08, ax_w - 0.06, f"{v}%", size=6.4, font=MONO, color=MUTED, h=0.16, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    data = [("T0", "Prompt only", 0, 0, 0), ("T1", "+ Retrieval", 38, 26, 52), ("T2", "+ Tools, vote", 82, 70, 92),
            ("T3", "+ Guardrail", 82, 70, 92), ("T4", "LoRA finetuned", 50, 36, 64)]
    slot = pw / 5; bwid = slot * 0.56
    for i, (k, nm, v, lo, hi) in enumerate(data):
        cx = px + slot * i + slot / 2
        c = GOLD if k in ("T2", "T3") else SLATE
        if v > 0:
            box(cx - bwid / 2, top + ph * (1 - v / 100), bwid, ph * v / 100, fill=c)
            vline(cx, top + ph * (1 - hi / 100), top + ph * (1 - lo / 100), color=INK, lw=1)
            for yy in (hi, lo):
                hline(cx - 0.06, top + ph * (1 - yy / 100), cx + 0.06, color=INK, lw=1)
        lab_y = top + ph * (1 - (hi if v else 0) / 100) - 0.2
        text(cx - slot / 2, lab_y, slot, f"{v}%", size=8.6, font=COND, bold=True, color=INK, h=0.18, align=PP_ALIGN.CENTER)
        text(cx - slot / 2, top + ph + 0.05, slot, k, size=7.6, font=MONO, bold=True, color=INK, h=0.16, align=PP_ALIGN.CENTER)
        text(cx - slot / 2 + 0.01, top + ph + 0.21, slot - 0.02, nm, size=6.4, color=MUTED, h=0.28, align=PP_ALIGN.CENTER, spacing=1.15)
    y = top + ph + 0.52
    y = para(x, y, [[B("The department is where the prompt alone fails. "), "T0 gets the issue right 56% of the time but never the department, guessing names such as “Public Works”. Retrieval raises department accuracy to 70%, and the codebook card to 82%."]], size=8.2) + 0.08
    y = para(x, y, [[B("Finetuning helps, but less than grounding. "), "With T0's prompt, T4 routes 50% correctly and is the fastest design (9.4 s per request). Most of its misses are plausible but wrong departments."]], size=8.2)
    return y

def pair_rows(x, y, rows, lw_=1.32):
    bx = x + lw_; bw = IW - lw_ - 0.36
    for nm, sub, a, b in rows:
        text(x, y, lw_ - 0.05, nm, size=7.6, bold=True, color=INK, h=0.17, spacing=1.15)
        text(x, y + 0.145, lw_ - 0.05, sub, size=6.4, color=MUTED, h=0.15, spacing=1.15)
        for j, (v, c) in enumerate([(a, INK), (b, GOLD)]):
            by = y + 0.02 + j * 0.13
            box(bx, by, bw, 0.1, fill=TRACK)
            if v: box(bx, by, bw * v / 100, 0.1, fill=c)
            text(bx + bw + 0.05, by - 0.03, 0.34, f"{v}%", size=6.8, font=MONO, color=TEXT, h=0.17, anchor=MSO_ANCHOR.MIDDLE)
        y += 0.28
    return y

def p_breakdown(x, y):
    box(x, y + 0.05, 0.12, 0.08, fill=INK); text(x + 0.17, y, 1.0, "T2 grounded", size=7, color=MUTED, h=0.18, anchor=MSO_ANCHOR.MIDDLE)
    box(x + 1.05, y + 0.05, 0.12, 0.08, fill=GOLD); text(x + 1.22, y, 1.0, "T4 finetuned", size=7, color=MUTED, h=0.18, anchor=MSO_ANCHOR.MIDDLE)
    y += 0.24
    text(x, y, IW, "ROUTED CORRECTLY, BY GOLD DOMAIN", size=6.4, font=MONO, bold=True, color=MUTED, charsp=0.5); y += 0.2
    y = pair_rows(x, y, [("Parks", "n = 14", 86, 86), ("Buildings", "n = 11", 82, 73), ("Waste", "n = 12", 100, 25), ("Streets", "n = 13", 62, 15)])
    y += 0.06
    text(x, y, IW, "ROUTED CORRECTLY, BY INPUT TYPE", size=6.4, font=MONO, bold=True, color=MUTED, charsp=0.5); y += 0.2
    y = pair_rows(x, y, [("Free-text complaint", "n = 13", 62, 15), ("Service request", "or example text, n = 25", 84, 80), ("Codebook row", "names the issue, n = 12", 100, 25)])
    y += 0.04
    y = para(x, y, "Free-text complaints are closest to what residents type, and T2 routes 62% of them correctly. Codebook rows name the issue, which makes them easy. T4 is weak on waste, where it had only 29 DEV rows.", size=7.6, color=MUTED)
    return y

y = panel(2, TOP, "05", "RESULT", "Grounding raised routing accuracy from 0% to 82%", p_result)
panel(2, y + GAP, "06", "BREAKDOWN", "Free-text complaints are the hardest case", p_breakdown, fill_to=BOT)

# ===== Column 4 =====
def p_safety(x, y):
    sets = [("Mahika (buildings)", 28, 100, 100), ("Rutomo (waste)", 30, 86, 94), ("Mingchin (parks)", 22, 42, 100)]
    for nm, n, blk, ben in sets:
        text(x, y, IW, [[{"text": nm, "bold": True, "color": INK}, {"text": f"   {n} probes", "color": MUTED, "size": 6.8}]], size=7.8, h=0.17)
        y += 0.19
        box(x, y, IW, 0.17, fill=TRACK)
        box(x, y, IW * blk / 100, 0.17, fill=INK)
        lab = f"{blk}% of attacks blocked"
        if blk >= 60:
            text(x + 0.06, y, IW * blk / 100 - 0.1, lab, size=6.8, bold=True, color=WHITE, h=0.17, anchor=MSO_ANCHOR.MIDDLE)
        else:
            text(x + IW * blk / 100 + 0.06, y, IW * (1 - blk / 100) - 0.1, lab, size=6.8, bold=True, color=INK, h=0.17, anchor=MSO_ANCHOR.MIDDLE)
        y += 0.19
        text(x, y, IW, f"Benign requests passed: {ben}%", size=6.8, color=MUTED, h=0.15); y += 0.24
    y += 0.0
    y = para(x, y, "Our guardrail is a keyword and pattern filter. It misses polite or newly phrased attacks, which make up most of Mingchin's probes. Without a guard, instructions planted in a corpus record worked for every member who tried them, so we treat the corpus as untrusted input.", size=7.6)
    return y

def p_lessons(x, y):
    items = [("A repetition penalty damaged the tickets.", "We set it to 1.15 to stop runaway output, but it also penalized prompt tokens the ticket has to copy, and T4 returned empty categories. Removing it raised T2 from 76% to 82%."),
             ("Our first finetune had too little signal.", "149 rows, one epoch and targets without a clarification question gave 0% routed correctly. With 269 rows, two epochs and card-based targets, T4 reached 50%."),
             ("Some inputs gave away the answer.", "A few member inputs stated where the request had been routed. We removed those sentences and rewrote codebook rows as resident reports.")]
    for i, (h, d) in enumerate(items):
        box(x, y, 0.24, 0.24, fill=INK, shape=MSO_SHAPE.OVAL)
        text(x, y, 0.24, str(i + 1), size=8, font=MONO, bold=True, color=GOLD, h=0.24, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        yy = y + text(x + 0.34, y + 0.02, IW - 0.34, h, size=8.2, bold=True, color=INK, spacing=1.2)
        yy = yy + 0.04 + text(x + 0.34, yy + 0.04, IW - 0.34, d, size=7.6, color=TEXT, spacing=1.25)
        y = yy + 0.12
    return y - 0.12

def p_limits(x, y):
    items = [("Small test set. ", "With 50 EVAL inputs, each accuracy can move about 13 points, so we rely only on large gaps."),
             ("Few realistic inputs. ", "Only 13 inputs are free text; on real resident wording we expect closer to 62% than 82%."),
             ("Resolution time is not an outcome. ", "Historical data cannot show that our API makes requests close faster.")]
    for h, d in items:
        box(x + 0.02, y + 0.065, 0.06, 0.06, fill=INK, shape=MSO_SHAPE.OVAL)
        y = para(x + 0.16, y, [[B(h), d]], size=7.6, w=IW - 0.16, spacing=1.25) + 0.07
    y += 0.03
    nh = 0.56
    box(x, y, IW, nh, fill=INK)
    text(x + 0.12, y + 0.07, IW - 0.24, [[{"text": "NEXT STEP", "font": MONO, "bold": True, "color": GOLD, "size": 6.6, "charsp": 0.8}],
                                        [{"text": "A pilot with City staff on real complaints, with staff confirming each routing decision.", "color": WHITE, "bold": True}]],
         size=8.2, h=nh - 0.1, spacing=1.2)
    return y + nh

y = panel(3, TOP, "07", "SAFETY", "The guardrail blocks only known attacks", p_safety)
y = panel(3, y + GAP, "08", "LESSONS", "What went wrong and what we changed", p_lessons)
panel(3, y + GAP, "09", "LIMITS", "Limitations and next step", p_limits, fill_to=BOT)

out = sys.argv[1] if len(sys.argv) > 1 else "poster_a3.pptx"
prs.save(out)
print("saved", out, f"col width {CW:.2f} in")
for w in WARN: print("WARN", w)
